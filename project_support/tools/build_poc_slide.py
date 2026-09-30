"""Build the editable PoC slide (16:9) with python-pptx: every text is a text box, every box a
shape, the KPI table a real table, icons and screenshots pictures. Coordinates are the pixel
positions of the HTML mock-up (1920x1080 at 144 px/in) so the two match."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

OUT_DIR = Path(r"D:\ICDCDT\project_support\output\poc_slide")
ICONS = OUT_DIR / "icons"
FONT = "맑은 고딕"

NAVY = RGBColor(0x0B, 0x1F, 0x5C)
NAVY2 = RGBColor(0x16, 0x3A, 0x8A)
PURPLE = RGBColor(0x5B, 0x3F, 0xA0)
TEAL = RGBColor(0x1D, 0x6B, 0x6B)
INK = RGBColor(0x1C, 0x1F, 0x2A)
MUTED = RGBColor(0x5C, 0x62, 0x70)
LINE = RGBColor(0xD9, 0xDD, 0xE8)
LINE_SOFT = RGBColor(0xE9, 0xEB, 0xF2)
BOX = RGBColor(0xF7, 0xF8, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
YELLOW = RGBColor(0xFF, 0xE2, 0x7A)
SKY = RGBColor(0x2B, 0x8F, 0xC7)
GREEN = RGBColor(0x2E, 0xA5, 0x6A)
RED = RGBColor(0xC8, 0x40, 0x2F)
BLACK = RGBColor(0x11, 0x11, 0x11)
GREEN_BG = RGBColor(0xED, 0xF8, 0xF1)
GREEN_LINE = RGBColor(0xBF, 0xE5, 0xCC)
GREEN_INK = RGBColor(0x2C, 0x4A, 0x38)
CHEV1 = RGBColor(0xCF, 0xD2, 0xDA)
CHEV2 = RGBColor(0xA9, 0xAD, 0xB8)
CHEV3 = RGBColor(0x6F, 0x74, 0x82)
FRAME = RGBColor(0x1A, 0x1D, 0x24)


def px(value):
    return Emu(int(round(value * 914400 / 144)))


def pt(px_value):
    return Pt(px_value / 2)


prs = Presentation()
prs.slide_width = px(1920)
prs.slide_height = px(1080)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def set_font(run, size_px, bold=False, color=INK, italic=False):
    font = run.font
    font.name = FONT
    font.size = pt(size_px)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text_box(x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, name=None):
    """paragraphs: list of (runs, options) where runs is a list of (text, size_px, bold, color)
    and options may hold align, space_after (px), line_spacing."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for runs, options in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = options.get("align", align)
        if "space_after" in options:
            p.space_after = pt(options["space_after"])
        if "space_before" in options:
            p.space_before = pt(options["space_before"])
        if "line_spacing" in options:
            p.line_spacing = options["line_spacing"]
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            set_font(r, size, bold, color)
    return tb


def para(text, size, bold=False, color=INK, **options):
    return ([(text, size, bold, color)], options)


def rect(x, y, w, h, fill=WHITE, line=None, line_px=2, radius=None, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if name:
        s.name = name
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = pt(line_px)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def rounded(x, y, w, h, fill=WHITE, line=LINE, line_px=2, radius_px=16, name=None):
    return rect(x, y, w, h, fill, line, line_px, radius=min(0.5, radius_px / min(w, h)), shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=name)


def circle(x, y, d, fill, name=None):
    return rect(x, y, d, d, fill, None, shape=MSO_SHAPE.OVAL, name=name)


def line(x1, y1, x2, y2, color=LINE_SOFT, width_px=1.5):
    # A thin filled rectangle: connectors pick up theme effects in some renderers.
    return rect(x1, y1 - width_px / 2, x2 - x1, width_px, color, None, name="separator")


def cell_bottom_border(cell, color="E6E9F2", width_emu=9525):
    tc_pr = cell._tc.get_or_add_tcPr()
    ln = tc_pr.makeelement(qn("a:lnB"), {"w": str(width_emu), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
    fill.append(clr)
    ln.append(fill)
    dash = ln.makeelement(qn("a:prstDash"), {"val": "solid"})
    ln.append(dash)
    tc_pr.insert(0, ln)


def picture(path, x, y, w, h, name=None):
    pic = slide.shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    if name:
        pic.name = name
    return pic


def heading(x, y, text, w=400):
    rect(x, y + 2, 7, 28, NAVY)
    text_box(x + 14, y, w, 32, [para(text, 24, True, NAVY)], anchor=MSO_ANCHOR.MIDDLE, name=f"heading {text}")


def chevron(x, y, w, h, fill, lines, color=INK):
    """Downward pentagon: a right-pointing home plate rotated 90°."""
    s = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, px(x + (w - h) / 2), px(y + (h - w) / 2), px(h), px(w))
    s.rotation = 90
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    s.adjustments[0] = 0.18
    s.text_frame.text = ""
    text_box(x, y, w, h - 22, [para(t, 22, True, color, align=PP_ALIGN.CENTER) for t in lines], anchor=MSO_ANCHOR.MIDDLE, name="chevron " + " ".join(lines))


# ---------- header ----------
text_box(190, 34, 900, 40, [para("PoC 추진 현황 및 주요 결과", 30, True, NAVY)], name="kicker")
text_box(44, 92, 1100, 84, [([("PoC 추진방안 ", 68, True, BLACK), ("| 건국대학교", 46, True, BLACK)], {})], anchor=MSO_ANCHOR.BOTTOM, name="title")
text_box(48, 212, 1100, 30, [para("ISDC 운용 디지털 트윈 기반 데이터 센터 운용 개념 검증", 22, True, INK)], name="subtitle")
text_box(1372, 108, 500, 50, [([("ISDC ODT v0.2", 17, True, NAVY), (" · 2026-09-08", 15, False, MUTED)], {"align": PP_ALIGN.RIGHT}),
                               para("SpaceTwin VVP · 8개 탭 운용 콘솔 · ICD 8종", 15, False, MUTED, align=PP_ALIGN.RIGHT)], name="version")

TOP = 262
BOX_TOP = TOP + 36

# ---------- left: concept ----------
heading(48, TOP, "개념 및 연계 구조")
rounded(48, BOX_TOP, 758, 668, WHITE, RGBColor(0xCF, 0xE3, 0xEE), name="left box")
row1, row2, row3 = BOX_TOP + 24, BOX_TOP + 24 + 214 + 14, BOX_TOP + 24 + 214 + 14 + 240 + 14
chevron(70, row1, 128, 214, CHEV1, ["PoC DT", "환경 정의", "범위"])
chevron(70, row2, 128, 240, CHEV2, ["ISDC", "ODT", "v0.2"])
chevron(70, row3, 128, 142, CHEV3, ["추진", "목표"], color=WHITE)

icons = [("folder", "데이터\n관리", "ICD-01 · DM-01~09", SKY), ("network", "데이터\n송/수신", "ICD-02 · DF-01~05", GREEN),
         ("document", "군집 노드\n운용개념", "ICD-03 · OR-01~04", PURPLE), ("target", "Target\nScenario", "ICD-06 · VF-01~03", TEAL)]
col_w = 560 / 4
for index, (icon, label, sub, color) in enumerate(icons):
    cx = 224 + index * col_w
    picture(ICONS / f"{icon}.png", cx + (col_w - 78) / 2, row1 + 6, 78, 78, name=f"icon {icon}")
    text_box(cx, row1 + 92, col_w, 60, [para(t, 21, True, color, align=PP_ALIGN.CENTER) for t in label.split("\n")], name=f"label {icon}")
    text_box(cx - 6, row1 + 156, col_w + 12, 20, [para(sub, 13, True, MUTED, align=PP_ALIGN.CENTER)], name=f"sub {icon}")

shot_w = (560 - 18) / 2
for index, (image, caption) in enumerate([("03_step2_fault.png", "DT 통합 운용 환경"), ("07_verdict_crop.png", "복구 판정 · KPI 비교")]):
    sx = 224 + index * (shot_w + 18)
    rect(sx, row2, shot_w, 190, FRAME, None, name=f"screen frame {index + 1}")
    picture(OUT_DIR / image, sx + 3, row2 + 3, shot_w - 6, 184, name=f"screenshot {index + 1}")
    rect(sx + shot_w / 2 - 60, row2 + 190, 120, 4, FRAME, None)
    text_box(sx, row2 + 200, shot_w, 30, [para(caption, 20, True, INK, align=PP_ALIGN.CENTER)], name=f"caption {index + 1}")

text_box(232, row3, 546, 142, [para("“ISDC ODT 기반 데이터 센터 운용 기술에 대한", 25, True, INK, align=PP_ALIGN.CENTER),
                               para("개념을 검증”", 25, True, INK, align=PP_ALIGN.CENTER, space_after=6),
                               para("링크 장애 → 상태 전달 → 우회 재전송 → 복구 판정", 21, True, NAVY, align=PP_ALIGN.CENTER)],
         anchor=MSO_ANCHOR.MIDDLE, name="goal")

# ---------- middle: scenario ----------
MX, MW = 838, 486
heading(MX, TOP, "검증 시나리오 · SDC_POC_01", 470)
rounded(MX, BOX_TOP, MW, 350, name="flow box")
text_box(MX + 18, BOX_TOP + 12, 300, 26, [para("OISL 주 링크 단절과 우회 복구", 17, True, NAVY)], name="flow title")
text_box(MX + 300, BOX_TOP + 16, 168, 20, [para("5단계 · 서버 SIM 시계 동기", 12, True, MUTED, align=PP_ALIGN.RIGHT)], name="flow title note")
text_box(MX + 18, BOX_TOP + 38, MW - 36, 34, [para("Walker Δ 4면 × 10기 · 550 km / 53° · 지상국 3곳 · 임무 3건(중계·관측 확정, 군집 갱신 계획)", 12, True, MUTED)], name="flow sub")
steps = [("1", "정상 운용", "N1→N2→N4 주 경로 2홉 전달 · 기준 KPI 표본", "T+0", NAVY, NAVY2),
         ("2", "Fault 주입", "편성된 주 링크 N1–N2에 link_loss 장애(300 s)", "T+20 s", RED, RED),
         ("3", "상태 전달", "DF-01 상태 갱신 → 패브릭 링크 사용 불가 판정", "T+24 s", NAVY, NAVY2),
         ("4", "우회 · 재전송", "장애 링크 제외 OR-01 재편성 · 우회 4홉 재전송", "T+30 s", NAVY, NAVY2),
         ("5", "복구 판정", "주 경로 복귀 · DM-02 서비스 · KPI 비교 판정(VF-03)", "T+435 s", GREEN, GREEN)]
for index, (num, title, desc, when, dot, when_color) in enumerate(steps):
    y = BOX_TOP + 78 + index * 54
    rounded(MX + 18, y, MW - 36, 48, BOX, LINE_SOFT, 1, 10, name=f"step {num} card")
    circle(MX + 28, y + 11, 26, dot, name=f"step {num} number")
    text_box(MX + 28, y + 11, 26, 26, [para(num, 13, True, WHITE, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    text_box(MX + 64, y + 5, 320, 40, [para(title, 14.5, True, INK), para(desc, 11.5, False, MUTED)], anchor=MSO_ANCHOR.MIDDLE, name=f"step {num} text")
    text_box(MX + MW - 18 - 80, y, 70, 48, [para(when, 13, True, when_color, align=PP_ALIGN.RIGHT)], anchor=MSO_ANCHOR.MIDDLE, name=f"step {num} time")

RY = BOX_TOP + 350 + 16
rounded(MX, RY, MW, 286, name="result box")
text_box(MX + 18, RY + 12, 200, 26, [para("복구 판정 결과", 17, True, NAVY)], name="result title")
text_box(MX + 200, RY + 16, 268, 20, [para("ICD-06 VF-03 · 3/3 항목 통과", 12, True, MUTED, align=PP_ALIGN.RIGHT)], name="result title note")
cards = [("도달 ✓", "중계 임무 완료 · 기한 여유 6,770 s"), ("자원 ✓", "N1 보관량 0 % · 데이터 안정성 75→77"), ("상태 ✓", "재수렴 1 s · 주 경로 복귀 · 보안 유지")]
card_w = (MW - 36 - 16) / 3
for index, (title, desc) in enumerate(cards):
    cx = MX + 18 + index * (card_w + 8)
    rounded(cx, RY + 42, card_w, 62, GREEN_BG, GREEN_LINE, 1, 10, name=f"verdict card {index + 1}")
    text_box(cx + 9, RY + 46, card_w - 18, 56, [para(title, 13, True, GREEN, space_after=1), para(desc, 11, False, GREEN_INK)], anchor=MSO_ANCHOR.MIDDLE)

rows = [["지표", "정상", "장애", "우회", "복구"],
        ["경로 홉 수", "2", "4", "4", "2"],
        ["편도 지연", "31.6 ms", "40.4 ms", "40.4 ms", "31.6 ms"],
        ["임무 경로 홉/종단 지연", "2 / 16 ms", "2 / 16 ms", "4 / 32 ms", "4 / 32 ms"],
        ["기한 여유", "7,003 s", "7,003 s", "6,770 s", "6,770 s"]]
highlight = {(1, 2): RED, (1, 4): GREEN, (2, 2): RED, (2, 4): GREEN, (3, 3): RED, (4, 4): GREEN}
table_shape = slide.shapes.add_table(len(rows), 5, px(MX + 18), px(RY + 114), px(MW - 36), px(110))
table_shape.name = "KPI table"
table = table_shape.table
tbl_pr = table._tbl.tblPr
tbl_pr.set("firstRow", "0")
tbl_pr.set("bandRow", "0")
style = tbl_pr.find(qn("a:tableStyleId"))
if style is None:
    style = tbl_pr.makeelement(qn("a:tableStyleId"), {})
    tbl_pr.append(style)
style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
widths = [166, 71, 71, 71, 71]
for column, width in zip(table.columns, widths):
    column.width = px(width)
for r, row in enumerate(rows):
    table.rows[r].height = px(22)
    for c, value in enumerate(row):
        cell = table.cell(r, c)
        cell.margin_left = cell.margin_right = px(5)
        cell.margin_top = cell.margin_bottom = px(1)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = BOX if r == 0 else WHITE
        cell_bottom_border(cell)
        tf = cell.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = value
        color = highlight.get((r, c), MUTED if r == 0 else INK)
        set_font(run, 11 if r == 0 else 12, r == 0 or c == 0 or (r, c) in highlight, color)
text_box(MX + 18, RY + 238, MW - 36, 40, [para("장애 주입 T+20 s → 대체 경로 확인 T+21 s(재수렴 1 s) → 재구성 요청 T+31 s → 장애 해제 T+320 s → 주 경로 복귀 T+435 s. 값은 SIM과 협력 SW 임시 구현의 대표 공학값.", 11, False, MUTED)], name="result note")

# ---------- right: items ----------
RX, RW = 1352, 520
heading(RX, TOP, "기능 검증 항목")
rounded(RX, BOX_TOP, RW, 380, name="items box")
items = [
    ("w_nodes", "위성 40기 배치 · 상태 동기화", "", "위성을 4개 궤도면에 띄우고 앞·뒤·좌·우 위성과 광 링크로 잇는다. 화면과 서버가 같은 배치를 본다."),
    ("w_wave", "링크 판정 · 경로 찾기", "", "매초 위성 위치와 링크 상태를 보내면 쓸 수 있는 링크를 가려내고 가장 빠른 길을 찾는다. 링크가 끊기면 돌아가는 길을 낸다."),
    ("w_power", "임무 배정 · 재배정", "", "어느 위성이 언제 무엇을 할지 정해 일정표에 올린다. 링크가 끊기면 그 링크를 빼고 새 경로로 다시 배정한다."),
    ("w_refresh", "데이터 보관 · 서비스", "", "위성마다 저장 공간을 등록하고 데이터를 어디에 몇 벌 둘지 정한다. 장애 중에도 데이터 요청이 처리되는지 본다."),
    ("w_monitor", "보안 · 복구 판정", "", "장애 중에도 보안 기준이 유지되는지 확인하고, 정상·장애·우회·복구 구간의 수치를 비교해 복구 성공 여부를 판정한다."),
]


def item_rows(box_top, box_h, entries, color, prefix):
    row_h = (box_h - 18) / len(entries)
    for index, (icon, title, tag, desc) in enumerate(entries):
        y = box_top + 10 + index * row_h
        circle(RX + 18, y + (row_h - 46) / 2, 46, color, name=f"{prefix} {index + 1} circle")
        picture(ICONS / f"{icon}.png", RX + 18 + 10, y + (row_h - 46) / 2 + 10, 26, 26, name=f"{prefix} {index + 1} icon")
        text_box(RX + 82, y + 4, RW - 100, row_h - 8, [([(title, 15.5, True, INK)] + ([("  " + tag, 12, True, color)] if tag else []), {"space_after": 1}), para(desc, 12, False, MUTED)],
                 anchor=MSO_ANCHOR.MIDDLE, name=f"{prefix} {index + 1} text")
        if index < len(entries) - 1:
            line(RX + 18, y + row_h, RX + RW - 18, y + row_h)


item_rows(BOX_TOP, 380, items, PURPLE, "item")
DY = BOX_TOP + 380 + 16
heading(RX, DY, "설계 반영 방향")
rounded(RX, DY + 36, RW, 236, name="design box")
designs = [("w_cube", "모듈 경계 확정", "DT ↔ 운용 SW", "DT는 상태 계산·소유, 운용 SW는 판단 · ICD 메시지로만 연결(구현/모의/계획 표기)"),
           ("w_plus", "협력 SW 교체 경로", "동일 ICD", "임시 구현 → 외부 모듈 주소 지정(--data-fabric-url 등)으로 무수정 교체"),
           ("w_chevrons", "후속 확장", "계획 메시지", "시험 결과 저장·조회(VF-04) · 번들 전달 상태(DF-06) · 군집 상태 보고(OR-05)")]
item_rows(DY + 36, 236, designs, TEAL, "design")

# ---------- banner ----------
BY = 1080 - 24 - 90
rounded(48, BY, 1824, 90, NAVY, None, 0, 18, name="banner")
picture(ICONS / "target_white.png", 78, BY + 12, 66, 66, name="banner icon")
text_box(166, BY, 1680, 90, [([("ISDC ODT 기반 ", 27, True, WHITE), ("OISL 링크 장애 → 상태 전달 → 우회 재전송 → 복구 판정", 27, True, YELLOW),
                                (" 전 과정을 DT와 협력 운용 SW의 ", 27, True, WHITE), ("ICD 연동", 27, True, YELLOW), ("으로 검증합니다.", 27, True, WHITE)], {})],
         anchor=MSO_ANCHOR.MIDDLE, name="banner text")

out = OUT_DIR / "poc_slide_v1.pptx"
prs.save(out)
print("saved", out, out.stat().st_size, "bytes;", len(slide.shapes), "shapes")
